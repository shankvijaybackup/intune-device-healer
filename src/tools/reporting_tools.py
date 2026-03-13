"""
Advanced Reporting Tools
"""

import os
from typing import Optional, Dict, Any, List
from datetime import datetime
import structlog
from docx import Document
from openpyxl import Workbook
from tools.monitoring import MonitoringTools

logger = structlog.get_logger()

class ReportingTools:
    """Tools for generating professional reports"""

    def __init__(self, authenticator, config):
        self.authenticator = authenticator
        self.config = config
        self.monitoring = MonitoringTools(authenticator, config)

    async def generate_fleet_health_report(self, format: str = "docx") -> Dict[str, Any]:
        """Generate a comprehensive fleet health report in professional format"""
        dashboard = await self.monitoring.get_fleet_health_dashboard()
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_dir = os.path.join(os.getcwd(), "reports")
        os.makedirs(report_dir, exist_ok=True)
        
        filename = f"fleet_health_report_{timestamp}.{format}"
        filepath = os.path.join(report_dir, filename)

        if format == "docx":
            self._create_docx_report(dashboard, filepath)
        elif format == "xlsx":
            self._create_xlsx_report(dashboard, filepath)
        else:
            return {"success": False, "error": f"Unsupported format: {format}"}

        return {
            "success": True, 
            "format": format,
            "filepath": filepath,
            "summary": dashboard.get("summary", {})
        }

    def _create_docx_report(self, data: Dict[str, Any], filepath: str):
        doc = Document()
        doc.add_heading('Fleet Health Report', 0)
        
        doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        
        # Summary Section
        doc.add_heading('Summary', level=1)
        summary = data.get("summary", {})
        doc.add_paragraph(f"Total Devices: {summary.get('total_devices', 'N/A')}")
        doc.add_paragraph(f"Healthy Devices: {summary.get('healthy_devices', 'N/A')}")
        doc.add_paragraph(f"Health Score: {summary.get('fleet_health_score', 'N/A')}/100")

        # Top Issues
        doc.add_heading('Top Issues', level=1)
        for issue in data.get("top_issues", []):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{issue.get('type', 'Unknown')}: ").bold = True
            p.add_run(f"{issue.get('count', 0)} devices impacted")

        doc.save(filepath)

    def _create_xlsx_report(self, data: Dict[str, Any], filepath: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Fleet Summary"
        
        ws['A1'] = "Metric"
        ws['B1'] = "Value"
        
        summary = data.get("summary", {})
        ws.append(["Total Devices", summary.get('total_devices')])
        ws.append(["Healthy Devices", summary.get('healthy_devices')])
        ws.append(["Fleet Health Score", summary.get('fleet_health_score')])
        
        # Issues sheet
        ws_issues = wb.create_sheet(title="Top Issues")
        ws_issues.append(["Issue Type", "Count", "Severity"])
        for issue in data.get("top_issues", []):
            ws_issues.append([issue.get('type'), issue.get('count'), issue.get('severity')])

        wb.save(filepath)
